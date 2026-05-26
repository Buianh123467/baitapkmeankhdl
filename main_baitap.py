import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from sklearn.cluster import KMeans
import openpyxl
from openpyxl.styles import Font

# =========================
# ĐỌC FILE EXCEL
# =========================
try:
    df = pd.read_excel("DiemTongHop.xlsx", header=None)
except:
    print("Không tìm thấy file DiemTongHop.xlsx")
    exit()

# =========================
# LẤY THÔNG TIN SINH VIÊN
# =========================
mssv = df.iloc[1, 3:].values
ten_sv = df.iloc[2, 3:].values

# LẤY BẢNG ĐIỂM
scores = df.iloc[4:56, 3:]

# CHUYỂN THÀNH SỐ
scores = scores.apply(pd.to_numeric, errors='coerce')

# XÓA CỘT RỖNG
scores = scores.dropna(axis=1, how='all')

# THAY Ô TRỐNG THÀNH 0
scores = scores.fillna(0)

# TÍNH GPA
gpa = scores.mean(axis=0)

# TẠO DATAFRAME
students = pd.DataFrame({
    "MSSV": mssv[:len(gpa)],
    "TenSV": ten_sv[:len(gpa)],
    "GPA": gpa.values
})

# =========================
# K-MEANS
# =========================
X = students[['GPA']]

kmeans = KMeans(
    n_clusters=3,
    random_state=42,
    n_init=10
)

students['Cluster'] = kmeans.fit_predict(X)

# SẮP XẾP CỤM
avg = students.groupby('Cluster')['GPA'].mean().sort_values()

mapping = {
    avg.index[0]: "Trung bình",
    avg.index[1]: "Khá",
    avg.index[2]: "Giỏi"
}

students['HocLuc'] = students['Cluster'].map(mapping)

# =========================
# IN KẾT QUẢ
# =========================
print(students[['TenSV', 'GPA', 'HocLuc']])

# =========================
# BIỂU ĐỒ PHÂN CỤM
# =========================
plt.figure(figsize=(10,5))

colors = {
    "Giỏi": "green",
    "Khá": "orange",
    "Trung bình": "red"
}

for hl in colors:
    temp = students[students['HocLuc'] == hl]

    plt.scatter(
        temp['GPA'],
        np.random.normal(0, 0.03, len(temp)),
        c=colors[hl],
        s=120,
        label=hl
    )

plt.title("Phân cụm sinh viên bằng K-Means")
plt.xlabel("GPA")
plt.legend()

plt.savefig("phan_cum.png")

# =========================
# BIỂU ĐỒ TRÒN
# =========================
plt.figure(figsize=(7,7))

counts = students['HocLuc'].value_counts()

plt.pie(
    counts,
    labels=counts.index,
    autopct='%1.1f%%'
)

plt.title("Tỷ lệ học lực")

plt.savefig("ty_le.png")

# =========================
# XUẤT EXCEL
# =========================
wb = openpyxl.Workbook()
ws = wb.active

headers = ["STT", "MSSV", "Tên Sinh Viên", "GPA", "Học Lực"]

for i, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=i, value=h)
    cell.font = Font(bold=True)

for idx, row in students.iterrows():

    ws.cell(idx+2, 1, idx+1)
    ws.cell(idx+2, 2, str(row['MSSV']))
    ws.cell(idx+2, 3, row['TenSV'])
    ws.cell(idx+2, 4, round(row['GPA'], 2))
    ws.cell(idx+2, 5, row['HocLuc'])

# chỉnh độ rộng cột
ws.column_dimensions['A'].width = 10
ws.column_dimensions['B'].width = 18
ws.column_dimensions['C'].width = 30
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 20

wb.save("KetQuaKMeans.xlsx")

print("\nĐã xuất file KetQuaKMeans.xlsx")
print("Đã lưu biểu đồ phan_cum.png")
print("Đã lưu biểu đồ ty_le.png")